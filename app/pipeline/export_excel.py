from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.models.schemas import ProcessingResult, ProcessingStatus

RECEIPTS_HEADERS = [
    "receipt_id",
    "date",
    "merchant",
    "category",
    "currency",
    "subtotal",
    "tax",
    "tip",
    "discount",
    "total",
    "payment_method",
    "confidence",
    "needs_review",
    "status",
    "source_file",
]

LINE_ITEM_HEADERS = [
    "receipt_id",
    "description",
    "quantity",
    "unit_price",
    "total_price",
    "confidence",
]

REVIEW_HEADERS = ["receipt_id", "issue", "field", "extracted_value", "source_file"]
LOG_HEADERS = [
    "receipt_id",
    "document_type",
    "route",
    "ocr_engine",
    "classifier_confidence",
    "ocr_confidence",
    "processing_time_ms",
    "status",
    "error",
]


def export_results(results: list[ProcessingResult], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(output_path) if output_path.exists() else Workbook()
    _ensure_sheet(workbook, "Receipts", RECEIPTS_HEADERS)
    _ensure_sheet(workbook, "Line Items", LINE_ITEM_HEADERS)
    _ensure_sheet(workbook, "Review Queue", REVIEW_HEADERS)
    _ensure_sheet(workbook, "Processing Logs", LOG_HEADERS)

    receipts = workbook["Receipts"]
    line_items = workbook["Line Items"]
    review = workbook["Review Queue"]
    logs = workbook["Processing Logs"]

    for result in results:
        extraction = result.extraction
        source = str(result.receipt_input.source_path)
        if extraction is not None:
            receipts.append(
                [
                    extraction.receipt_id,
                    extraction.date.isoformat() if extraction.date else None,
                    extraction.merchant_name,
                    extraction.category,
                    extraction.currency,
                    extraction.subtotal,
                    extraction.tax,
                    extraction.tip,
                    extraction.discount,
                    extraction.total,
                    extraction.payment_method,
                    extraction.confidence,
                    extraction.needs_review,
                    result.status.value,
                    source,
                ]
            )
            for item in extraction.line_items:
                line_items.append(
                    [
                        extraction.receipt_id,
                        item.description,
                        item.quantity,
                        item.unit_price,
                        item.total_price,
                        item.confidence,
                    ]
                )
            for warning in extraction.warnings:
                review.append([extraction.receipt_id, warning, None, None, source])
        elif result.status == ProcessingStatus.FAILED:
            review.append([result.receipt_input.receipt_id, result.error, None, None, source])

        logs.append(
            [
                result.receipt_input.receipt_id,
                result.classification.document_type,
                result.classification.route.value,
                result.classification.ocr_engine,
                result.classification.confidence,
                result.document.confidence if result.document else None,
                result.processing_time_ms,
                result.status.value,
                result.error,
            ]
        )

    _autosize(workbook["Receipts"])
    _autosize(workbook["Line Items"])
    _autosize(workbook["Review Queue"])
    _autosize(workbook["Processing Logs"])
    workbook.save(output_path)


def _ensure_sheet(workbook: Workbook, name: str, headers: list[str]) -> None:
    if name in workbook.sheetnames:
        sheet = workbook[name]
        if sheet.max_row == 0:
            sheet.append(headers)
        return
    if workbook.active.title == "Sheet" and workbook.active.max_row == 1 and workbook.active.max_column == 1:
        sheet = workbook.active
        sheet.title = name
    else:
        sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8F1FF")


def _autosize(sheet: Worksheet) -> None:
    for column_cells in sheet.columns:
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 12), 48)
